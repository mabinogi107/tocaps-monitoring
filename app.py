from flask import Flask, request, jsonify, send_from_directory
import openpyxl, os, re
from collections import defaultdict, Counter

app = Flask(__name__, static_folder='static')

# ── 분류 기준
현장_패턴 = [
    '보령발전본부','서대문은평지사','부천삼정','광양 바이오매스',
    '유진마포','가마지천','중곡-','한국건강관리협회',
    '산포처리구역','일산수질복원센터','국회대로','이태원111',
    '푸른솔 GC','신축공사','리모델링','정비공사','정비사업',
    '구축공사','조성공사','개선공사','건립 ','건설공사','시설개량','지하차도',
    '뉴스테이','주상복합','아파트','오피스텔','물류센터','데이터센터',
    'EPC','현장','공사현장'
]
건설본사_부서 = {
    '공사관리팀','CM팀','사업관리팀','P.I TFT','PI TFT',
    '프리콘팀','프리콘팀(프로젝트A)','프리콘팀(프로젝트B)',
    '건설솔루션팀','외주구매팀','공공사업실',
    '환경에너지사업실','공사관리실(공사관리Part)',
    '건축사업실','사업전략담당','건설사업본부','기술서비스담당',
    'DC프리콘 TFT','개발운영팀','DC프리콘TFT'
}
회전기_부서 = {
    '회전기영업팀','회전기영업팀(영업1Part)','회전기영업팀(영업2Part)',
    '회전기사업본부','예산공장(생산관리Part)','예산공장(공장지원Part)',
    '예산공장(품질보증Part)','예산공장(회전기안전Part)','예산공장장',
    'VPC(설계Part)','VPC','회전기사업/VPC(관리Part)'
}

CAT_LABELS = ['본사', '건설(본사)', '건설(현장)', '회전기사업본부']

def get_cat(dept):
    if not dept: return '미분류'
    d = str(dept)
    if d in 회전기_부서: return '회전기사업본부'
    if d in 건설본사_부서: return '건설(본사)'
    for p in 현장_패턴:
        if p in d: return '건설(현장)'
    return '본사'

def read_sheet(ws, col_grpname, has_no=False):
    rows = []
    for r in range(4, 99999):
        e = ws.cell(r,5).value
        f = ws.cell(r,6).value
        if e is None and f is None:
            break
        rows.append({
            'e': str(e or '').strip(),
            'f': str(f or '').strip(),
            'g': str(ws.cell(r,7).value or '').strip(),
            'h': str(ws.cell(r,8).value or '').strip(),
            'i': str(ws.cell(r,9).value or '').strip(),
            'name': str(ws.cell(r,col_grpname).value or '').strip(),
            'c': str(ws.cell(r,3).value or '').strip(),
            'n': str(ws.cell(r,14).value or '') if has_no else '',
            'o': str(ws.cell(r,15).value or '') if has_no else '',
            'biго': str(ws.cell(r,16).value or '') if has_no else '',
        })
    return rows

def analyze(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)

    sheets = wb.sheetnames
    # 전월/당월 시트 탐색
    prev_ws = curr_ws = None
    for s in sheets:
        if '전월' in s: prev_ws = wb[s]
        if '당월' in s: curr_ws = wb[s]
    if prev_ws is None or curr_ws is None:
        raise ValueError(f'전월/당월 시트를 찾을 수 없습니다. 시트 목록: {sheets}')

    # 헤더에서 모니터링 기간 찾기
    period = ''
    for r in range(1, 5):
        for c in range(1, 10):
            v = str(curr_ws.cell(r,c).value or '')
            if '모니터링' in v or '기간' in v:
                period = v.strip()
                break

    prev_rows = read_sheet(prev_ws, col_grpname=10, has_no=True)
    curr_rows = read_sheet(curr_ws, col_grpname=12)

    prev_g = [r for r in prev_rows if r['c'] == '건설']
    curr_g = [r for r in curr_rows if r['c'] == '건설']

    def make_key(r):
        return (r['e'], r['i'], r['g'])

    prev_by_key = {}
    for r in prev_g:
        k = make_key(r)
        if k not in prev_by_key:
            prev_by_key[k] = r

    curr_by_key = {}
    for r in curr_g:
        k = make_key(r)
        if k not in curr_by_key:
            curr_by_key[k] = r

    prev_set = set(prev_by_key.keys())
    curr_set = set(curr_by_key.keys())

    removed_keys = prev_set - curr_set
    added_keys   = curr_set - prev_set

    removed_list = [prev_by_key[k] for k in removed_keys]
    added_list   = [curr_by_key[k] for k in added_keys]

    prev_cnt = Counter(get_cat(r['h']) for r in prev_g)
    curr_cnt = Counter(get_cat(r['h']) for r in curr_g)

    removed_by_cat = defaultdict(list)
    for r in removed_list:
        removed_by_cat[get_cat(r['h'])].append({
            'name': r['f'], 'dept': r['h'], 'code': r['i'], 'grp': r['name']
        })

    added_by_cat = defaultdict(list)
    for r in added_list:
        added_by_cat[get_cat(r['h'])].append({
            'name': r['f'], 'dept': r['h'], 'code': r['i'], 'grp': r['name']
        })

    # N/O 누락
    n_miss = [{'name':r['f'],'dept':r['h'],'code':r['i'],'cat':get_cat(r['h'])} for r in prev_g if r['n']=='누락']
    o_miss = [{'name':r['f'],'dept':r['h'],'code':r['i'],'cat':get_cat(r['h'])} for r in prev_g if r['o']=='누락']

    # 결과 조합
    cats_result = []
    for cat in CAT_LABELS:
        p = prev_cnt[cat]
        c = curr_cnt[cat]
        rm = sorted(removed_by_cat[cat], key=lambda x: x['name'])
        ad = sorted(added_by_cat[cat],   key=lambda x: x['name'])
        cats_result.append({
            'cat': cat,
            'prev': p,
            'curr': c,
            'diff': c - p,
            'removed': rm,
            'added': ad,
            'check': p - len(rm) + len(ad) == c,
        })

    # 전체 소속별 집계
    소속별 = Counter(r['c'] for r in curr_rows)

    return {
        'period': period,
        'sheets': sheets,
        'prev_total': len(prev_g),
        'curr_total': len(curr_g),
        'cats': cats_result,
        'n_miss': n_miss,
        'o_miss': o_miss,
        '소속별': dict(소속별),
    }

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/analyze', methods=['POST'])
def analyze_route():
    if 'file' not in request.files:
        return jsonify({'error': '파일이 없습니다'}), 400
    f = request.files['file']
    tmp = os.path.join(os.environ.get('TEMP', '.'), 'tocaps_upload.xlsx')
    f.save(tmp)
    try:
        result = analyze(tmp)
        return jsonify(result)
    except Exception as ex:
        return jsonify({'error': str(ex)}), 500
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)

if __name__ == '__main__':
    os.makedirs('static', exist_ok=True)
    app.run(debug=False, port=5050)
